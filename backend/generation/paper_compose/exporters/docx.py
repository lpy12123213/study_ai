
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import List, Optional

from backend.core.logging_utils import get_logger

logger = get_logger(__name__)

AI_SYNTHESIS_REVIEW_NOTE = "※ 本题答案由 AI 生成，请复核。"


def _is_ai_synthesis(q: dict) -> bool:
    return str(q.get("answer_source") or q.get("answerSource") or "").strip() == "ai_synthesis"


try:
    from docx import Document  # type: ignore[import-not-found]
except ImportError:  # pragma: no cover
    Document = None  # type: ignore[assignment]

def render_paper_docx_bytes(
    paper: dict,
    *,
    include_stem: bool = False,
    include_answer: bool = False,
    include_analysis: bool = False,
) -> bytes:
    # Prefer pandoc when available: it converts LaTeX math into real Word equations (OMML),
    # which python-docx cannot easily do.
    engine = str(os.getenv("PAPER_EXPORT_DOCX_ENGINE") or "").strip().lower()

    def _pandoc_available() -> Optional[str]:
        return shutil.which("pandoc")

    def _render_exam_markdown() -> str:
        name = str(paper.get("paper_name") or paper.get("name") or "试卷").strip()
        pid = str(paper.get("paper_id") or paper.get("id") or "").strip()
        created = str(paper.get("created_at") or paper.get("createdAt") or "").strip()
        questions = paper.get("questions") if isinstance(paper.get("questions"), list) else []

        lines: List[str] = [f"# {name}", ""]
        meta: List[str] = []
        if pid:
            meta.append(f"试卷ID：{pid}")
        if created:
            meta.append(f"创建时间：{created}")
        if meta:
            lines.append("**" + " | ".join(meta) + "**")
            lines.append("")

        # Group by question type (preserve first-seen ordering).
        groups: List[tuple[str, List[dict]]] = []
        buckets: dict[str, List[dict]] = {}
        for q in questions:
            if not isinstance(q, dict):
                continue
            qtype = str(q.get("type") or q.get("question_type") or "").strip() or "其他题"
            if qtype not in buckets:
                buckets[qtype] = []
                groups.append((qtype, buckets[qtype]))
            buckets[qtype].append(q)

        answer_lines: List[str] = []
        qnum = 0
        for group_name, qs in groups:
            lines.append(f"## {group_name}")
            lines.append("")
            for q in qs:
                if not isinstance(q, dict):
                    continue
                qnum += 1
                qid = str(q.get("question_id") or q.get("questionId") or "").strip()
                lines.append(f"### {qnum}. {qid}".strip())
                lines.append("")

                if include_stem:
                    stem = str(q.get("stem") or "").strip()
                    if stem:
                        lines.append(stem)
                        lines.append("")

                diagrams = q.get("diagrams")
                if isinstance(diagrams, list) and diagrams:
                    for d in diagrams[:3]:
                        if not isinstance(d, dict):
                            continue
                        filename = str(d.get("filename") or "").strip()
                        url = str(d.get("url") or "").strip()
                        if not filename and "/api/media/generated/" in url:
                            filename = (
                                url.split("/api/media/generated/", 1)[-1].split("?", 1)[0].strip().strip("/")
                            )
                        if filename:
                            lines.append(f"![]({filename})")
                            lines.append("")

                if include_answer or include_analysis:
                    ans = str(q.get("answer") or "").strip() if include_answer else ""
                    ana = str(q.get("analysis") or "").strip() if include_analysis else ""
                    if ans or ana:
                        answer_lines.append(f"#### {qnum}.")
                        if ans:
                            answer_lines.append(f"**答案：** {ans}")
                            if _is_ai_synthesis(q):
                                answer_lines.append(f"> {AI_SYNTHESIS_REVIEW_NOTE}")
                        if ana:
                            answer_lines.append(f"**解析：** {ana}")
                        answer_lines.append("")

        if answer_lines:
            lines.append("---")
            lines.append("")
            lines.append("## 参考答案")
            lines.append("")
            lines.extend(answer_lines)

        return "\n".join(lines).rstrip() + "\n"

    def _copy_generated_assets(*, markdown: str, tmp_dir: Path) -> None:
        repo_root = Path(__file__).resolve().parents[3]
        generated_dir = (repo_root / ".local" / "media" / "generated").resolve()
        if not generated_dir.exists():
            return

        # Copy only referenced local filenames in markdown image tags.
        for m in re.finditer(r"!\\[[^\\]]*\\]\\(([^\\)]+)\\)", markdown or ""):
            raw = str(m.group(1) or "").strip()
            if not raw or ":" in raw:
                continue
            filename = raw.split("?", 1)[0].strip().replace("\\", "/").split("/")[-1]
            if not filename:
                continue
            src = (generated_dir / filename).resolve()
            try:
                src.relative_to(generated_dir)
            except ValueError:
                continue
            if not src.exists() or not src.is_file():
                continue
            dst = (tmp_dir / filename).resolve()
            if dst.exists():
                continue
            try:
                shutil.copyfile(str(src), str(dst))
            except OSError:
                continue

    def _add_page_number_footer(doc) -> None:  # type: ignore[no-untyped-def]
        try:
            from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-not-found]
            from docx.oxml import OxmlElement  # type: ignore[import-not-found]
            from docx.oxml.ns import qn  # type: ignore[import-not-found]
        except ImportError:
            return

        try:
            section = doc.sections[0]
            footer = section.footer
            para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            fld = OxmlElement("w:fldSimple")
            fld.set(qn("w:instr"), "PAGE")
            run._r.append(fld)  # type: ignore[attr-defined]
        except (AttributeError, IndexError, TypeError, ValueError):
            return

    def _render_via_pandoc(*, markdown: str, timeout_s: float = 60.0) -> Optional[bytes]:
        pandoc = _pandoc_available()
        if not pandoc:
            return None

        timeout_s = float(timeout_s or 60.0)
        timeout_s = max(5.0, min(timeout_s, 60.0 * 10.0))

        with tempfile.TemporaryDirectory(prefix="paper_docx_") as tmp:
            tmp_dir = Path(tmp)
            md_path = tmp_dir / "paper.md"
            docx_path = tmp_dir / "paper.docx"
            md_path.write_text(markdown, encoding="utf-8")
            _copy_generated_assets(markdown=markdown, tmp_dir=tmp_dir)

            cmd = [
                pandoc,
                "--from",
                "markdown+tex_math_dollars+tex_math_single_backslash",
                "--to",
                "docx",
                "-o",
                str(docx_path),
                str(md_path),
            ]
            try:
                proc = subprocess.run(
                    cmd,
                    cwd=str(tmp_dir),
                    capture_output=True,
                    text=True,
                    timeout=timeout_s,
                    check=False,
                )
            except (OSError, subprocess.SubprocessError):
                return None

            if proc.returncode != 0 or not docx_path.exists():
                return None

            try:
                raw = docx_path.read_bytes()
            except OSError:
                return None

            if Document is None:
                return raw

            # Post-process footer page numbers when python-docx is available.
            try:
                doc = Document(io.BytesIO(raw))
                _add_page_number_footer(doc)
                out_buf = io.BytesIO()
                doc.save(out_buf)
                return out_buf.getvalue()
            except Exception:
                logger.warning("paper_export_docx_footer_postprocess_failed", exc_info=True)
                return raw

    if engine in {"", "pandoc", "auto"}:
        md = _render_exam_markdown()
        via = _render_via_pandoc(markdown=md, timeout_s=float(os.getenv("PAPER_EXPORT_PANDOC_TIMEOUT_S") or "60"))
        if via:
            return via

    if Document is None:
        raise ValueError("docx_not_available")

    doc = Document()
    name = str(paper.get("paper_name") or paper.get("name") or "试卷").strip()
    pid = str(paper.get("paper_id") or paper.get("id") or "").strip()
    created = str(paper.get("created_at") or paper.get("createdAt") or "").strip()
    questions = paper.get("questions") if isinstance(paper.get("questions"), list) else []

    # Title + meta
    doc.add_heading(name, level=0)
    meta_parts: List[str] = []
    if pid:
        meta_parts.append(f"试卷ID：{pid}")
    if created:
        meta_parts.append(f"创建时间：{created}")
    if meta_parts:
        doc.add_paragraph(" | ".join(meta_parts))

    # Group by question type (preserve order).
    groups: List[tuple[str, List[dict]]] = []
    buckets: dict[str, List[dict]] = {}
    for q in questions:
        if not isinstance(q, dict):
            continue
        qtype = str(q.get("type") or q.get("question_type") or "").strip() or "其他题"
        if qtype not in buckets:
            buckets[qtype] = []
            groups.append((qtype, buckets[qtype]))
        buckets[qtype].append(q)

    repo_root = Path(__file__).resolve().parents[3]
    generated_dir = (repo_root / ".local" / "media" / "generated").resolve()

    try:
        from docx.shared import Inches  # type: ignore[import-not-found]
    except ImportError:  # pragma: no cover
        Inches = None  # type: ignore[assignment]

    def _add_paragraph_lines(text: str) -> None:
        t = str(text or "")
        if not t.strip():
            doc.add_paragraph("")
            return
        for line in t.splitlines():
            doc.add_paragraph(line)

    qnum = 0
    for group_name, qs in groups:
        doc.add_heading(group_name, level=1)
        for q in qs:
            if not isinstance(q, dict):
                continue
            qnum += 1
            qid = str(q.get("question_id") or q.get("questionId") or "").strip()
            doc.add_heading(f"{qnum}. {qid}".strip(), level=2)

            if include_stem:
                _add_paragraph_lines(str(q.get("stem") or "").strip())

            diagrams = q.get("diagrams")
            if isinstance(diagrams, list) and diagrams and generated_dir.exists():
                for d in diagrams[:3]:
                    if not isinstance(d, dict):
                        continue
                    filename = str(d.get("filename") or "").strip()
                    url = str(d.get("url") or "").strip()
                    if not filename and "/api/media/generated/" in url:
                        filename = url.split("/api/media/generated/", 1)[-1].split("?", 1)[0].strip().strip("/")
                    if not filename:
                        continue
                    path = (generated_dir / filename).resolve()
                    try:
                        path.relative_to(generated_dir)
                    except ValueError:
                        continue
                    if not path.exists() or not path.is_file():
                        continue
                    if path.suffix.lower() == ".svg":
                        continue
                    try:
                        if Inches is not None:
                            doc.add_picture(str(path), width=Inches(5.6))
                        else:
                            doc.add_picture(str(path))
                    except Exception:
                        logger.warning("paper_export_docx_image_failed", extra={"path": str(path)}, exc_info=True)
                        continue
                    caption = str(d.get("caption") or "").strip()
                    if caption:
                        doc.add_paragraph(caption)

            if include_answer:
                ans = str(q.get("answer") or "").strip()
                if ans:
                    doc.add_paragraph("答案：")
                    _add_paragraph_lines(ans)
                    if _is_ai_synthesis(q):
                        doc.add_paragraph(AI_SYNTHESIS_REVIEW_NOTE)

            if include_analysis:
                ana = str(q.get("analysis") or "").strip()
                if ana:
                    doc.add_paragraph("解析：")
                    _add_paragraph_lines(ana)

    _add_page_number_footer(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
