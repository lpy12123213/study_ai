from __future__ import annotations

import asyncio
import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from backend.database.models import get_question_cache
from backend.media.generated import default_generated_media_ttl_s, publish_generated_bytes, publish_generated_text
from backend.paper_compose.exam_templates import format_answer_key_section, format_exam_header, get_exam_preamble

try:
    from docx import Document  # type: ignore[import-not-found]
except Exception:  # pragma: no cover
    Document = None  # type: ignore[assignment]


_TEX_ESCAPE_REPL = {
    "\\": r"\textbackslash{}",
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
    "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}",
}
_TEX_ESCAPE_PATTERN = re.compile(r"[\\&%$#_{}~^]")
_TEX_MATH_SPAN_PATTERN = re.compile(r"(\\\((?:.|\n)*?\\\)|\\\[(?:.|\n)*?\\\])", re.DOTALL)


def _escape_plain_tex(text: str) -> str:
    # IMPORTANT: use a single-pass regexp replacement so the replacement
    # strings (e.g. \textbackslash{}) won't be re-escaped again.
    t = str(text or "")
    return _TEX_ESCAPE_PATTERN.sub(lambda m: _TEX_ESCAPE_REPL.get(m.group(0), m.group(0)), t)


def _smart_tex_escape(text: str) -> str:
    """Escape TeX special chars outside math spans \\(...\\) and \\[...\\]."""

    raw = str(text or "")
    if not raw:
        return ""

    parts = _TEX_MATH_SPAN_PATTERN.split(raw)
    if len(parts) <= 1:
        return _escape_plain_tex(raw)

    out: List[str] = []
    for idx, part in enumerate(parts):
        if not part:
            continue
        if idx % 2 == 1 and (part.startswith("\\(") and part.endswith("\\)")):
            out.append(part)
            continue
        if idx % 2 == 1 and (part.startswith("\\[") and part.endswith("\\]")):
            out.append(part)
            continue
        out.append(_escape_plain_tex(part))
    return "".join(out)


def render_paper_markdown(
    paper: dict,
    *,
    include_stem: bool = False,
    include_answer: bool = False,
    include_analysis: bool = False,
) -> str:
    name = str(paper.get("paper_name") or paper.get("name") or "试卷").strip()
    pid = paper.get("paper_id") or paper.get("id") or ""
    created = str(paper.get("created_at") or paper.get("createdAt") or "").strip()
    questions = paper.get("questions") if isinstance(paper.get("questions"), list) else []

    lines = [f"# {name}", ""]
    meta = []
    if pid:
        meta.append(f"> 试卷 ID：{pid}")
    if created:
        meta.append(f"> 创建时间：{created}")
    if meta:
        lines.extend(meta)
        lines.append("")

    lines.append("| 题号 | 题型 | 难度 | 知识点 | 题目ID | 来源 |")
    lines.append("| --- | --- | --- | --- | --- | --- |")
    for q in questions:
        if not isinstance(q, dict):
            continue
        order = q.get("order") or q.get("question_order") or ""
        qtype = str(q.get("type") or q.get("question_type") or "").strip()
        diff = str(q.get("difficulty") or "").strip()
        kp = str(q.get("knowledge_point") or q.get("knowledgePoint") or "").strip()
        qid = str(q.get("question_id") or q.get("questionId") or "").strip()
        src = str(q.get("source_url") or q.get("sourceUrl") or "").strip()
        lines.append(f"| {order} | {qtype} | {diff} | {kp} | {qid} | {src} |")

    if include_stem or include_answer or include_analysis:
        lines.append("")
        lines.append("---")
        lines.append("")
        lines.append("## 题目内容（本地快照）")
        lines.append("")

        for q in questions:
            if not isinstance(q, dict):
                continue
            qid = str(q.get("question_id") or q.get("questionId") or "").strip()
            order = q.get("order") or q.get("question_order") or ""
            lines.append(f"### {order}. {qid}")
            lines.append("")

            if include_stem:
                stem = str(q.get("stem") or "").strip()
                if stem:
                    lines.append("**题干：**")
                    lines.append("")
                    lines.append(stem)
                    lines.append("")

            if include_answer:
                ans = str(q.get("answer") or "").strip()
                if ans:
                    lines.append("**答案：**")
                    lines.append("")
                    lines.append(ans)
                    lines.append("")

            if include_analysis:
                ana = str(q.get("analysis") or "").strip()
                if ana:
                    lines.append("**解析：**")
                    lines.append("")
                    lines.append(ana)
                    lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def render_paper_docx_bytes(
    paper: dict,
    *,
    include_stem: bool = False,
    include_answer: bool = False,
    include_analysis: bool = False,
) -> bytes:
    # Prefer pandoc when available: it converts LaTeX math into real Word equations (OMML),
    # which python-docx cannot easily do.
    engine = str(os.getenv("PAPER_EXPORT_DOCX_ENGINE") or "").strip().lower()

    def _pandoc_available() -> Optional[str]:
        return shutil.which("pandoc")

    def _render_exam_markdown() -> str:
        name = str(paper.get("paper_name") or paper.get("name") or "试卷").strip()
        pid = str(paper.get("paper_id") or paper.get("id") or "").strip()
        created = str(paper.get("created_at") or paper.get("createdAt") or "").strip()
        questions = paper.get("questions") if isinstance(paper.get("questions"), list) else []

        lines: List[str] = [f"# {name}", ""]
        meta: List[str] = []
        if pid:
            meta.append(f"试卷ID：{pid}")
        if created:
            meta.append(f"创建时间：{created}")
        if meta:
            lines.append("**" + " | ".join(meta) + "**")
            lines.append("")

        # Group by question type (preserve first-seen ordering).
        groups: List[tuple[str, List[dict]]] = []
        buckets: dict[str, List[dict]] = {}
        for q in questions:
            if not isinstance(q, dict):
                continue
            qtype = str(q.get("type") or q.get("question_type") or "").strip() or "其他题"
            if qtype not in buckets:
                buckets[qtype] = []
                groups.append((qtype, buckets[qtype]))
            buckets[qtype].append(q)

        answer_lines: List[str] = []
        qnum = 0
        for group_name, qs in groups:
            lines.append(f"## {group_name}")
            lines.append("")
            for q in qs:
                if not isinstance(q, dict):
                    continue
                qnum += 1
                qid = str(q.get("question_id") or q.get("questionId") or "").strip()
                lines.append(f"### {qnum}. {qid}".strip())
                lines.append("")

                if include_stem:
                    stem = str(q.get("stem") or "").strip()
                    if stem:
                        lines.append(stem)
                        lines.append("")

                diagrams = q.get("diagrams")
                if isinstance(diagrams, list) and diagrams:
                    for d in diagrams[:3]:
                        if not isinstance(d, dict):
                            continue
                        filename = str(d.get("filename") or "").strip()
                        url = str(d.get("url") or "").strip()
                        if not filename and "/api/media/generated/" in url:
                            filename = (
                                url.split("/api/media/generated/", 1)[-1].split("?", 1)[0].strip().strip("/")
                            )
                        if filename:
                            lines.append(f"![]({filename})")
                            lines.append("")

                if include_answer or include_analysis:
                    ans = str(q.get("answer") or "").strip() if include_answer else ""
                    ana = str(q.get("analysis") or "").strip() if include_analysis else ""
                    if ans or ana:
                        answer_lines.append(f"#### {qnum}.")
                        if ans:
                            answer_lines.append(f"**答案：** {ans}")
                        if ana:
                            answer_lines.append(f"**解析：** {ana}")
                        answer_lines.append("")

        if answer_lines:
            lines.append("---")
            lines.append("")
            lines.append("## 参考答案")
            lines.append("")
            lines.extend(answer_lines)

        return "\n".join(lines).rstrip() + "\n"

    def _copy_generated_assets(*, markdown: str, tmp_dir: Path) -> None:
        repo_root = Path(__file__).resolve().parents[2]
        generated_dir = (repo_root / ".local" / "media" / "generated").resolve()
        if not generated_dir.exists():
            return

        # Copy only referenced local filenames in markdown image tags.
        for m in re.finditer(r"!\\[[^\\]]*\\]\\(([^\\)]+)\\)", markdown or ""):
            raw = str(m.group(1) or "").strip()
            if not raw or ":" in raw:
                continue
            filename = raw.split("?", 1)[0].strip().replace("\\", "/").split("/")[-1]
            if not filename:
                continue
            src = (generated_dir / filename).resolve()
            try:
                src.relative_to(generated_dir)
            except Exception:
                continue
            if not src.exists() or not src.is_file():
                continue
            dst = (tmp_dir / filename).resolve()
            if dst.exists():
                continue
            try:
                shutil.copyfile(str(src), str(dst))
            except Exception:
                continue

    def _add_page_number_footer(doc) -> None:  # type: ignore[no-untyped-def]
        try:
            from docx.oxml import OxmlElement  # type: ignore[import-not-found]
            from docx.oxml.ns import qn  # type: ignore[import-not-found]
            from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-not-found]
        except Exception:
            return

        try:
            section = doc.sections[0]
            footer = section.footer
            para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            fld = OxmlElement("w:fldSimple")
            fld.set(qn("w:instr"), "PAGE")
            run._r.append(fld)  # type: ignore[attr-defined]
        except Exception:
            return

    def _render_via_pandoc(*, markdown: str, timeout_s: float = 60.0) -> Optional[bytes]:
        pandoc = _pandoc_available()
        if not pandoc:
            return None

        timeout_s = float(timeout_s or 60.0)
        timeout_s = max(5.0, min(timeout_s, 60.0 * 10.0))

        with tempfile.TemporaryDirectory(prefix="paper_docx_") as tmp:
            tmp_dir = Path(tmp)
            md_path = tmp_dir / "paper.md"
            docx_path = tmp_dir / "paper.docx"
            md_path.write_text(markdown, encoding="utf-8")
            _copy_generated_assets(markdown=markdown, tmp_dir=tmp_dir)

            cmd = [
                pandoc,
                "--from",
                "markdown+tex_math_dollars+tex_math_single_backslash",
                "--to",
                "docx",
                "-o",
                str(docx_path),
                str(md_path),
            ]
            try:
                proc = subprocess.run(
                    cmd,
                    cwd=str(tmp_dir),
                    capture_output=True,
                    text=True,
                    timeout=timeout_s,
                    check=False,
                )
            except Exception:
                return None

            if proc.returncode != 0 or not docx_path.exists():
                return None

            try:
                raw = docx_path.read_bytes()
            except Exception:
                return None

            if Document is None:
                return raw

            # Post-process footer page numbers when python-docx is available.
            try:
                doc = Document(io.BytesIO(raw))
                _add_page_number_footer(doc)
                out_buf = io.BytesIO()
                doc.save(out_buf)
                return out_buf.getvalue()
            except Exception:
                return raw

    if engine in {"", "pandoc", "auto"}:
        md = _render_exam_markdown()
        via = _render_via_pandoc(markdown=md, timeout_s=float(os.getenv("PAPER_EXPORT_PANDOC_TIMEOUT_S") or "60"))
        if via:
            return via

    if Document is None:
        raise ValueError("docx_not_available")

    doc = Document()
    name = str(paper.get("paper_name") or paper.get("name") or "试卷").strip()
    pid = str(paper.get("paper_id") or paper.get("id") or "").strip()
    created = str(paper.get("created_at") or paper.get("createdAt") or "").strip()
    questions = paper.get("questions") if isinstance(paper.get("questions"), list) else []

    # Title + meta
    doc.add_heading(name, level=0)
    meta_parts: List[str] = []
    if pid:
        meta_parts.append(f"试卷ID：{pid}")
    if created:
        meta_parts.append(f"创建时间：{created}")
    if meta_parts:
        doc.add_paragraph(" | ".join(meta_parts))

    # Group by question type (preserve order).
    groups: List[tuple[str, List[dict]]] = []
    buckets: dict[str, List[dict]] = {}
    for q in questions:
        if not isinstance(q, dict):
            continue
        qtype = str(q.get("type") or q.get("question_type") or "").strip() or "其他题"
        if qtype not in buckets:
            buckets[qtype] = []
            groups.append((qtype, buckets[qtype]))
        buckets[qtype].append(q)

    repo_root = Path(__file__).resolve().parents[2]
    generated_dir = (repo_root / ".local" / "media" / "generated").resolve()

    try:
        from docx.shared import Inches  # type: ignore[import-not-found]
    except Exception:  # pragma: no cover
        Inches = None  # type: ignore[assignment]

    def _add_paragraph_lines(text: str) -> None:
        t = str(text or "")
        if not t.strip():
            doc.add_paragraph("")
            return
        for line in t.splitlines():
            doc.add_paragraph(line)

    qnum = 0
    for group_name, qs in groups:
        doc.add_heading(group_name, level=1)
        for q in qs:
            if not isinstance(q, dict):
                continue
            qnum += 1
            qid = str(q.get("question_id") or q.get("questionId") or "").strip()
            doc.add_heading(f"{qnum}. {qid}".strip(), level=2)

            if include_stem:
                _add_paragraph_lines(str(q.get("stem") or "").strip())

            diagrams = q.get("diagrams")
            if isinstance(diagrams, list) and diagrams and generated_dir.exists():
                for d in diagrams[:3]:
                    if not isinstance(d, dict):
                        continue
                    filename = str(d.get("filename") or "").strip()
                    url = str(d.get("url") or "").strip()
                    if not filename and "/api/media/generated/" in url:
                        filename = url.split("/api/media/generated/", 1)[-1].split("?", 1)[0].strip().strip("/")
                    if not filename:
                        continue
                    path = (generated_dir / filename).resolve()
                    try:
                        path.relative_to(generated_dir)
                    except Exception:
                        continue
                    if not path.exists() or not path.is_file():
                        continue
                    if path.suffix.lower() == ".svg":
                        continue
                    try:
                        if Inches is not None:
                            doc.add_picture(str(path), width=Inches(5.6))
                        else:
                            doc.add_picture(str(path))
                    except Exception:
                        continue
                    caption = str(d.get("caption") or "").strip()
                    if caption:
                        doc.add_paragraph(caption)

            if include_answer:
                ans = str(q.get("answer") or "").strip()
                if ans:
                    doc.add_paragraph("答案：")
                    _add_paragraph_lines(ans)

            if include_analysis:
                ana = str(q.get("analysis") or "").strip()
                if ana:
                    doc.add_paragraph("解析：")
                    _add_paragraph_lines(ana)

    _add_page_number_footer(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_paper_latex(
    paper: dict,
    *,
    include_stem: bool = False,
    include_answer: bool = False,
    include_analysis: bool = False,
) -> str:
    name = str(paper.get("paper_name") or paper.get("name") or "试卷").strip()
    pid = str(paper.get("paper_id") or paper.get("id") or "").strip()
    created = str(paper.get("created_at") or paper.get("createdAt") or "").strip()
    questions = paper.get("questions") if isinstance(paper.get("questions"), list) else []

    def _generated_filename_from_url(url: str) -> str:
        raw = str(url or "").strip()
        if not raw:
            return ""
        raw = raw.split("?", 1)[0]
        token = "/api/media/generated/"
        if token in raw:
            raw = raw.split(token, 1)[-1]
        raw = raw.strip().strip("/").strip()
        # Only allow basenames to prevent path traversal in the compiler stage.
        raw = raw.replace("\\", "/").split("/")[-1]
        return raw

    # Standard exam layout (header -> grouped questions -> optional answer key).
    notes = [
        "请将所有答案写在答题纸上，写在本试卷上无效。",
        "作答前请仔细审题，注意步骤与表达规范。",
    ]
    if pid:
        notes.append(f"试卷ID：{_smart_tex_escape(pid)}")
    if created:
        notes.append(f"创建时间：{_smart_tex_escape(created)}")

    lines: List[str] = [get_exam_preamble(margin_cm=2.0), r"\begin{document}"]
    lines.append(
        format_exam_header(
            title=_smart_tex_escape(name),
            subject=_smart_tex_escape(str(paper.get("subject") or "").strip()),
            time_limit=_smart_tex_escape(str(paper.get("time_limit") or "").strip()),
            total_points=_smart_tex_escape(str(paper.get("total_points") or "").strip()),
            notes=notes,
        ).rstrip()
    )

    # Group by question type (preserve first-seen ordering).
    groups: List[tuple[str, List[dict]]] = []
    groups_by_name: dict[str, List[dict]] = {}
    for q in questions:
        if not isinstance(q, dict):
            continue
        qtype = str(q.get("type") or q.get("question_type") or "").strip() or "其他题"
        if qtype not in groups_by_name:
            groups_by_name[qtype] = []
            groups.append((qtype, groups_by_name[qtype]))
        groups_by_name[qtype].append(q)

    if not groups:
        lines.append(r"\section*{（无题目）}")
        lines.append(r"\end{document}")
        return "\n".join(lines).rstrip() + "\n"

    qnum = 0
    answers: List[dict] = []

    for group_name, qs in groups:
        lines.append(rf"\section*{{{_smart_tex_escape(group_name)}}}")
        lines.append("")
        for q in qs:
            if not isinstance(q, dict):
                continue
            qnum += 1
            qid = str(q.get("question_id") or q.get("questionId") or "").strip()
            points = q.get("points") if q.get("points") is not None else q.get("score")
            points_text = ""
            try:
                if points is not None and str(points).strip():
                    points_text = f"（{_smart_tex_escape(str(points).strip())}分）"
            except Exception:
                points_text = ""

            stem = str(q.get("stem") or "").strip()
            if not include_stem and not stem:
                stem_tex = rf"\textit{{题目ID：{_smart_tex_escape(qid) or 'unknown'}}}"
            elif include_stem:
                stem_tex = _smart_tex_escape(stem).replace("\n", r"\\")
            else:
                # include_stem=false but stem exists: keep it out for safety.
                stem_tex = rf"\textit{{题目ID：{_smart_tex_escape(qid) or 'unknown'}}}"

            lines.append(rf"\textbf{{{qnum}.}} {stem_tex} {points_text}".rstrip())

            # Optional diagrams field: emit local includegraphics (compiler copies assets).
            diagrams = q.get("diagrams")
            if isinstance(diagrams, list) and diagrams:
                for d in diagrams[:3]:
                    if not isinstance(d, dict):
                        continue
                    filename = str(d.get("filename") or "").strip() or _generated_filename_from_url(str(d.get("url") or ""))
                    if not filename:
                        continue
                    lines.append(r"\begin{center}")
                    lines.append(rf"\includegraphics[width=0.72\linewidth]{{{_smart_tex_escape(filename)}}}")
                    cap = str(d.get("caption") or "").strip()
                    if cap:
                        lines.append(rf"\\[0.2em]{{\small {_smart_tex_escape(cap)}}}")
                    lines.append(r"\end{center}")

            lines.append("")

            if include_answer or include_analysis:
                ans = str(q.get("answer") or "").strip() if include_answer else ""
                ana = str(q.get("analysis") or "").strip() if include_analysis else ""
                answer_tex = _smart_tex_escape(ans).replace("\n", r"\\") if ans else ""
                analysis_tex = _smart_tex_escape(ana).replace("\n", r"\\") if ana else ""
                if answer_tex or analysis_tex:
                    answers.append({"number": qnum, "answer_tex": answer_tex, "analysis_tex": analysis_tex})

    if answers:
        lines.append(format_answer_key_section(answers=answers).rstrip())

    lines.append(r"\end{document}")
    return "\n".join(lines).rstrip() + "\n"


def _find_latex_engine() -> Optional[str]:
    configured = str(os.getenv("PAPER_EXPORT_LATEX_ENGINE") or os.getenv("LATEX_ENGINE") or "").strip()
    if configured:
        return shutil.which(configured) or configured
    return shutil.which("xelatex") or shutil.which("pdflatex")


def _resolve_paper_export_latex_timeout_s(timeout_s: Optional[float]) -> float:
    raw = str(os.getenv("PAPER_EXPORT_LATEX_TIMEOUT_S") or "").strip()
    if timeout_s is None or float(timeout_s or 0) <= 0:
        # Default: keep it short to avoid accidental long hangs in API exports.
        # Users can override via env when they have a full TeX distribution and want longer runs.
        try:
            timeout_s = float(raw) if raw else 30.0
        except Exception:
            timeout_s = 30.0

    timeout_s = float(timeout_s or 30.0)
    return max(5.0, min(timeout_s, 60.0 * 20.0))


def compile_latex_to_pdf(*, tex: str, timeout_s: Optional[float] = None) -> Tuple[Optional[bytes], str]:
    """Compile LaTeX to PDF and return (pdf_bytes, log_text)."""

    engine = _find_latex_engine()
    if not engine:
        return None, "latex_engine_not_found"

    timeout_s = _resolve_paper_export_latex_timeout_s(timeout_s)

    def _copy_includegraphics_assets(*, tex_text: str, tmp_dir: Path) -> int:
        # Copy referenced images from `.local/media/generated/` into the build dir,
        # so \includegraphics{sha.png} can resolve.
        repo_root = Path(__file__).resolve().parents[2]
        generated_dir = (repo_root / ".local" / "media" / "generated").resolve()
        if not generated_dir.exists():
            return 0

        pattern = re.compile(r"\\includegraphics(?:\\[[^\\]]*\\])?\\{([^\\}]+)\\}")
        copied = 0
        for m in pattern.finditer(tex_text or ""):
            raw = str(m.group(1) or "").strip()
            if not raw:
                continue
            raw = raw.split("?", 1)[0].strip()
            # Only allow basenames to avoid escaping the temp dir.
            filename = raw.replace("\\", "/").split("/")[-1]
            if not filename or any(ch in filename for ch in [":", "\x00"]):
                continue

            # If extension omitted, try common ones.
            candidates = [filename]
            if "." not in filename:
                candidates = [f"{filename}{ext}" for ext in [".png", ".jpg", ".jpeg", ".pdf", ".svg"]]

            for cand in candidates:
                src = (generated_dir / cand).resolve()
                try:
                    src.relative_to(generated_dir)
                except Exception:
                    continue
                if not src.exists() or not src.is_file():
                    continue
                dst = (tmp_dir / cand).resolve()
                if dst.exists():
                    break
                try:
                    shutil.copyfile(str(src), str(dst))
                    copied += 1
                except Exception:
                    pass
                break
        return copied

    with tempfile.TemporaryDirectory(prefix="paper_export_") as tmp:
        tmp_dir = Path(tmp)
        tex_path = tmp_dir / "paper.tex"
        tex_path.write_text(tex, encoding="utf-8")

        _copy_includegraphics_assets(tex_text=tex, tmp_dir=tmp_dir)

        cmd = [
            engine,
            "-interaction=nonstopmode",
            "-halt-on-error",
            "-file-line-error",
            str(tex_path),
        ]
        last_log = ""
        for _ in range(2):
            try:
                proc = subprocess.run(
                    cmd,
                    cwd=str(tmp_dir),
                    capture_output=True,
                    text=True,
                    timeout=timeout_s,
                    check=False,
                )
            except subprocess.TimeoutExpired:
                return None, "latex_compile_timeout"
            except Exception as exc:
                return None, f"latex_compile_failed: {exc}"

            last_log = (proc.stdout or "") + "\n" + (proc.stderr or "")
            if proc.returncode != 0:
                break

        pdf_path = tmp_dir / "paper.pdf"
        if not pdf_path.exists():
            return None, last_log[-8000:]

        return pdf_path.read_bytes(), last_log[-8000:]


async def compile_latex_to_pdf_async(*, tex: str, timeout_s: Optional[float] = None) -> Tuple[Optional[bytes], str]:
    """Async wrapper around the sync compiler to avoid blocking the event loop."""

    timeout_s = _resolve_paper_export_latex_timeout_s(timeout_s)

    try:
        # We still keep a wall-clock timeout guard to prevent thread pool stalls.
        return await asyncio.wait_for(
            asyncio.to_thread(compile_latex_to_pdf, tex=tex, timeout_s=timeout_s),
            timeout=timeout_s + 5.0,
        )
    except asyncio.TimeoutError:
        return None, "latex_compile_timeout"


async def export_paper(
    paper: dict,
    *,
    user_id: str,
    fmt: str,
    include_stem: bool = False,
    include_answer: bool = False,
    include_analysis: bool = False,
) -> Dict[str, Any]:
    """Export paper into `.local/media/generated` and return URLs."""

    hydrated_paper = dict(paper or {})
    questions = hydrated_paper.get("questions") if isinstance(hydrated_paper.get("questions"), list) else []
    if (include_stem or include_answer or include_analysis) and questions:
        qids = []
        for q in questions:
            if not isinstance(q, dict):
                continue
            qid = str(q.get("question_id") or q.get("questionId") or "").strip()
            if qid:
                qids.append(qid)
        cache = await get_question_cache(question_ids=qids)
        for q in questions:
            if not isinstance(q, dict):
                continue
            qid = str(q.get("question_id") or q.get("questionId") or "").strip()
            rec = cache.get(qid) if qid else None
            if not isinstance(rec, dict):
                continue
            if include_stem and not str(q.get("stem") or "").strip():
                q["stem"] = str(rec.get("stem") or "").strip()
            if include_answer and not str(q.get("answer") or "").strip():
                q["answer"] = str(rec.get("answer") or "").strip()
            if include_analysis and not str(q.get("analysis") or "").strip():
                q["analysis"] = str(rec.get("analysis") or "").strip()
            if not str(q.get("source_url") or "").strip():
                q["source_url"] = str(rec.get("source_url") or "").strip()

    fmt_norm = (fmt or "").strip().lower()
    if fmt_norm in {"md", "markdown"}:
        md = render_paper_markdown(
            hydrated_paper,
            include_stem=include_stem,
            include_answer=include_answer,
            include_analysis=include_analysis,
        )
        out = await publish_generated_text(
            md,
            user_id=user_id,
            ext=".md",
            file_type="md",
            mime_type="text/markdown; charset=utf-8",
            ttl_s=default_generated_media_ttl_s(),
        )
        return {"format": "markdown", **out}

    if fmt_norm in {"tex", "latex"}:
        tex = render_paper_latex(
            hydrated_paper,
            include_stem=include_stem,
            include_answer=include_answer,
            include_analysis=include_analysis,
        )
        out = await publish_generated_text(
            tex,
            user_id=user_id,
            ext=".tex",
            file_type="tex",
            mime_type="application/x-tex; charset=utf-8",
            ttl_s=default_generated_media_ttl_s(),
        )
        return {"format": "latex", **out}

    if fmt_norm == "pdf":
        tex = render_paper_latex(
            hydrated_paper,
            include_stem=include_stem,
            include_answer=include_answer,
            include_analysis=include_analysis,
        )
        tex_out = await publish_generated_text(
            tex,
            user_id=user_id,
            ext=".tex",
            file_type="tex",
            mime_type="application/x-tex; charset=utf-8",
            ttl_s=default_generated_media_ttl_s(),
        )
        pdf_bytes, log_text = await compile_latex_to_pdf_async(tex=tex)
        if not pdf_bytes:
            return {
                "format": "pdf",
                "success": False,
                "error": "pdf_compile_failed",
                "log": log_text,
                "tex_url": tex_out["url"],
                "tex_filename": tex_out["filename"],
            }
        pdf_out = await publish_generated_bytes(
            pdf_bytes,
            user_id=user_id,
            ext=".pdf",
            file_type="pdf",
            mime_type="application/pdf",
            ttl_s=default_generated_media_ttl_s(),
        )
        return {
            "format": "pdf",
            "success": True,
            "pdf_url": pdf_out["url"],
            "pdf_filename": pdf_out["filename"],
            "tex_url": tex_out["url"],
            "tex_filename": tex_out["filename"],
        }

    if fmt_norm in {"docx", "word"}:
        docx_bytes = render_paper_docx_bytes(
            hydrated_paper,
            include_stem=include_stem,
            include_answer=include_answer,
            include_analysis=include_analysis,
        )
        out = await publish_generated_bytes(
            docx_bytes,
            user_id=user_id,
            ext=".docx",
            file_type="docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ttl_s=default_generated_media_ttl_s(),
        )
        return {"format": "docx", "success": True, **out}

    raise ValueError("unsupported_format")
